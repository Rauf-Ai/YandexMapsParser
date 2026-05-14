"""
Yandex Maps Parser
Extracts company data from Yandex Maps and exports to Excel.

Strategy:
  1. Yandex Search Maps API  — batch search (name, phone, address, coords, category)
  2. Per-org detail request  — full card (website, socials, rating, reviews)
"""

import argparse
import json
import logging
import random
import re
import sys
import time
from datetime import date
from urllib.parse import quote, urlencode

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# HTTP session
# ---------------------------------------------------------------------------

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept": "application/json, text/javascript, */*; q=0.01",
    "Referer": "https://yandex.ru/maps/",
}

SESSION = requests.Session()
SESSION.headers.update(HEADERS)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SEARCH_API     = "https://search-maps.yandex.ru/v1/"
SEARCH_API_KEY = "dda3ddba-c9ea-4ead-9010-f43fbc15c6e3"

SOCIAL_DOMAINS = (
    "vk.com", "vkontakte.ru",
    "instagram.com",
    "t.me", "telegram.me", "telegram.org",
    "ok.ru",
    "facebook.com", "fb.com",
    "youtube.com",
    "tiktok.com",
    "twitter.com", "x.com",
    "whatsapp.com",
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _sleep(lo: float = 0.8, hi: float = 2.0):
    time.sleep(random.uniform(lo, hi))


def _fmt_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("8") and len(digits) == 11:
        digits = "7" + digits[1:]
    if digits.startswith("7") and len(digits) == 11:
        return f"+7 ({digits[1:4]}) {digits[4:7]}-{digits[7:9]}-{digits[9:11]}"
    return raw.strip()


def _clean_site(url: str) -> str:
    url = re.sub(r"^https?://", "", url.strip())
    return url.rstrip("/")


def _is_social(url: str) -> bool:
    return any(d in url for d in SOCIAL_DOMAINS)


def _dedup_key(company: dict) -> tuple:
    return (company.get("name", "").lower(), company.get("address", "").lower())


def _extract_oid(uri: str) -> str | None:
    """Extract org ID from ymapsbm1://org?oid=12345"""
    m = re.search(r"oid=(\d+)", uri or "")
    return m.group(1) if m else None

# ---------------------------------------------------------------------------
# Step 1 — Search API (batch, basic fields)
# ---------------------------------------------------------------------------

def _search_page(query: str, skip: int) -> list[dict]:
    """Return list of raw feature dicts from Yandex Search Maps API."""
    params = {
        "apikey": SEARCH_API_KEY,
        "text": query,
        "lang": "ru_RU",
        "type": "biz",
        "results": 10,
        "skip": skip,
    }
    try:
        resp = SESSION.get(SEARCH_API, params=params, timeout=15)
        resp.raise_for_status()
        return resp.json().get("features", [])
    except Exception as exc:
        log.warning("Search API skip=%d: %s", skip, exc)
        return []


def _parse_feature(feat: dict) -> dict:
    """Parse a GeoJSON feature into a company dict (basic fields only)."""
    props = feat.get("properties", {})
    geo   = feat.get("geometry", {})

    name = props.get("name", "—")
    name = re.sub(r'^["\'«»]+|["\'»«]+$', "", name).strip() or "—"

    coords = geo.get("coordinates", [])
    lon = round(coords[0], 6) if len(coords) > 0 and coords[0] is not None else "—"
    lat = round(coords[1], 6) if len(coords) > 1 and coords[1] is not None else "—"

    meta     = props.get("CompanyMetaData", {})
    address  = meta.get("address", "—")
    cats     = meta.get("Categories", [])
    category = cats[0].get("name", "—") if cats else "—"
    oid      = _extract_oid(props.get("uri", "")) or meta.get("id", "")

    phones = [
        _fmt_phone(p.get("formatted", ""))
        for p in meta.get("Phones", [])
        if p.get("formatted")
    ]
    phone_str = ", ".join(phones) if phones else "—"

    # --- Rating (correct API paths) ---
    rating_obj = meta.get("rating") or {}
    # API returns {"ratings": 4.5, "reviews": 150}  or  {"score": 4.5, "count": 150}
    rating_raw  = (rating_obj.get("ratings") or rating_obj.get("score")
                   or rating_obj.get("value"))
    reviews_raw = (rating_obj.get("reviews") or rating_obj.get("count") or 0)

    try:
        rating = round(float(rating_raw), 1)
    except (TypeError, ValueError):
        rating = "—"
    try:
        reviews = int(reviews_raw)
    except (TypeError, ValueError):
        reviews = 0

    # --- Website & socials (from API Links) ---
    site    = meta.get("url", "") or ""          # sometimes a top-level field
    socials : list[str] = []

    for link in meta.get("Links", []):
        href = (link.get("href") or link.get("url") or "").strip()
        if not href:
            continue
        if _is_social(href):
            socials.append(_clean_site(href))
        elif not site:
            site = href

    site = _clean_site(site) if site else "—"

    return {
        "oid":      oid,
        "name":     name,
        "phone":    phone_str,
        "site":     site,
        "social":   ", ".join(socials) if socials else "—",
        "address":  address,
        "lat":      lat,
        "lon":      lon,
        "category": category,
        "rating":   rating,
        "reviews":  reviews,
        "has_site": "Да" if site != "—" else "Нет",
    }

# ---------------------------------------------------------------------------
# Step 2 — Per-org enrichment (full card: site, socials, rating, reviews)
# ---------------------------------------------------------------------------

def _enrich(company: dict) -> dict:
    """
    Fetch the org's Yandex Maps page and extract missing fields.
    Yandex embeds JSON state in the HTML that contains full business details.
    """
    oid = company.get("oid", "")
    if not oid:
        return company

    needs_site    = company["site"]    == "—"
    needs_social  = company["social"]  == "—"
    needs_rating  = company["rating"]  == "—"
    needs_reviews = company["reviews"] == 0

    if not any([needs_site, needs_social, needs_rating, needs_reviews]):
        return company  # already complete

    url = f"https://yandex.ru/maps/org/{oid}/"
    try:
        resp = SESSION.get(url, timeout=15)
        if resp.status_code != 200:
            return company
        html = resp.text
    except Exception as exc:
        log.debug("Enrich %s failed: %s", oid, exc)
        return company

    # --- Try to parse embedded JSON blobs ---
    # Yandex Maps embeds full org data in <script> tags as JSON
    json_blobs: list[dict] = []

    # Pattern 1: window.__data = {...}
    for m in re.finditer(r'window\.__(?:data|reduxState|REDUX_STATE__)\s*=\s*(\{.{50,})', html):
        try:
            raw = _balanced_json(m.group(1))
            json_blobs.append(json.loads(raw))
        except Exception:
            pass

    # Pattern 2: <script type="application/json">...</script>
    for tag in re.findall(r'<script[^>]*type="application/json"[^>]*>(.*?)</script>', html, re.DOTALL):
        try:
            json_blobs.append(json.loads(tag))
        except Exception:
            pass

    # Collect all string values from JSON blobs to find URLs
    flat_text = " ".join(_flatten_strings(b) for b in json_blobs) if json_blobs else html

    # --- Extract rating ---
    if needs_rating:
        for pat in [
            r'"rating"\s*:\s*\{[^}]*?"value"\s*:\s*([\d.]+)',
            r'"ratingValue"\s*:\s*([\d.]+)',
            r'"averageRating"\s*:\s*([\d.]+)',
            r'"score"\s*:\s*([\d.]+)',
        ]:
            m = re.search(pat, flat_text)
            if m:
                try:
                    company["rating"] = round(float(m.group(1)), 1)
                    needs_rating = False
                    break
                except ValueError:
                    pass

    # --- Extract review count ---
    if needs_reviews:
        for pat in [
            r'"reviewCount"\s*:\s*(\d+)',
            r'"reviews"\s*:\s*(\d+)',
            r'"ratingsCount"\s*:\s*(\d+)',
            r'"count"\s*:\s*(\d+)',
        ]:
            m = re.search(pat, flat_text)
            if m:
                try:
                    v = int(m.group(1))
                    if v > 0:
                        company["reviews"] = v
                        needs_reviews = False
                        break
                except ValueError:
                    pass

    # --- Extract website ---
    if needs_site:
        for pat in [
            r'"siteUrl"\s*:\s*"(https?://[^"]{4,})"',
            r'"url"\s*:\s*"(https?://(?!yandex\.|maps\.)[^"]{4,})"',
            r'"website"\s*:\s*"(https?://[^"]{4,})"',
        ]:
            m = re.search(pat, flat_text)
            if m:
                candidate = m.group(1)
                if not _is_social(candidate) and "yandex" not in candidate:
                    company["site"]     = _clean_site(candidate)
                    company["has_site"] = "Да"
                    needs_site = False
                    break

    # --- Extract socials ---
    if needs_social:
        found: list[str] = []
        for pat in [r'"(https?://(?:' + '|'.join(re.escape(d) for d in SOCIAL_DOMAINS) + r')[^"]*)"']:
            for m in re.finditer(pat, flat_text):
                link = _clean_site(m.group(1))
                if link not in found:
                    found.append(link)
        if found:
            company["social"] = ", ".join(found)

    # --- Fallback: parse visible HTML for rating / review badge ---
    if needs_rating or needs_reviews:
        soup = BeautifulSoup(html, "html.parser")
        if needs_rating:
            el = (soup.select_one(".business-rating-badge-view__rating") or
                  soup.select_one("[class*='rating__value']") or
                  soup.select_one("[itemprop='ratingValue']"))
            if el:
                try:
                    company["rating"] = round(float(
                        el.get("content") or el.get_text(strip=True).replace(",", ".")
                    ), 1)
                except ValueError:
                    pass
        if needs_reviews:
            el = (soup.select_one(".business-rating-badge-view__count") or
                  soup.select_one("[class*='rating__count']") or
                  soup.select_one("[itemprop='reviewCount']"))
            if el:
                try:
                    v = int(re.sub(r"\D", "", el.get("content") or el.get_text()))
                    if v > 0:
                        company["reviews"] = v
                except ValueError:
                    pass

    return company


def _balanced_json(s: str, max_len: int = 200_000) -> str:
    """Extract a balanced {...} JSON object from the start of s."""
    depth = 0
    in_str = False
    escape = False
    for i, ch in enumerate(s[:max_len]):
        if escape:
            escape = False
            continue
        if ch == "\\" and in_str:
            escape = True
            continue
        if ch == '"':
            in_str = not in_str
            continue
        if in_str:
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return s[:i + 1]
    return s


def _flatten_strings(obj, acc: list | None = None) -> str:
    """Recursively collect all string values from a nested dict/list."""
    if acc is None:
        acc = []
    if isinstance(obj, str):
        acc.append(obj)
    elif isinstance(obj, dict):
        for v in obj.values():
            _flatten_strings(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            _flatten_strings(v, acc)
    return " ".join(acc)

# ---------------------------------------------------------------------------
# Fallback: HTML search scraper
# ---------------------------------------------------------------------------

def _scrape_page(query: str, page: int) -> list[dict]:
    """HTML fallback when Search API returns nothing."""
    url = f"https://yandex.ru/maps/?text={quote(query)}&page={page}"
    try:
        resp = SESSION.get(url, timeout=15)
        resp.raise_for_status()
    except Exception as exc:
        log.warning("HTML page %d failed: %s", page, exc)
        return []

    soup = BeautifulSoup(resp.text, "html.parser")
    items = soup.select("li.search-list-item")
    companies = []
    for item in items:
        def txt(sel):
            el = item.select_one(sel)
            return el.get_text(strip=True) if el else "—"

        name    = txt(".search-business-snippet-view__title")
        address = txt(".search-business-snippet-view__address")
        category= txt(".search-business-snippet-view__category")
        try:
            rating = round(float(
                txt(".business-rating-badge-view__rating").replace(",", ".")
            ), 1)
        except ValueError:
            rating = "—"
        try:
            reviews = int(re.sub(r"\D", "", txt(".business-rating-badge-view__count")))
        except ValueError:
            reviews = 0

        # Try to extract oid from data attributes
        oid = ""
        for attr in ("data-permalink", "data-uri"):
            val = item.get(attr, "")
            m = re.search(r"(\d{10,})", val)
            if m:
                oid = m.group(1)
                break

        companies.append({
            "oid": oid, "name": name, "phone": "—",
            "site": "—", "social": "—", "address": address,
            "lat": "—", "lon": "—", "category": category,
            "rating": rating, "reviews": reviews, "has_site": "Нет",
        })
    return companies

# ---------------------------------------------------------------------------
# Main collection loop
# ---------------------------------------------------------------------------

def collect(query: str, max_companies: int = 100,
            emit=None) -> list[dict]:
    """
    Collect companies for query.
    emit(event, **kw) — optional callback for progress events (used by web UI).
    """
    def _emit(event, **kw):
        if emit:
            emit(event, **kw)

    companies: list[dict] = []
    seen: set[tuple] = set()
    skip = 0
    page = 1
    empty_streak = 0

    while len(companies) < max_companies:
        _emit("progress",
              found=len(companies), total=max_companies,
              message=f"Поиск результатов {skip + 1}–{skip + 10}…")

        feats = _search_page(query, skip)
        batch = [_parse_feature(f) for f in feats]

        if not batch:
            log.warning("API skip=%d пустой, HTML fallback…", skip)
            _emit("progress", found=len(companies), total=max_companies,
                  message=f"API не ответил, HTML fallback (стр. {page})…")
            batch = _scrape_page(query, page)

        if not batch:
            empty_streak += 1
            if empty_streak >= 3:
                _emit("progress", found=len(companies), total=max_companies,
                      message="Результаты закончились.")
                break
        else:
            empty_streak = 0

        for c in batch:
            key = _dedup_key(c)
            if key in seen:
                continue
            seen.add(key)

            # Enrich with per-org detail page
            _emit("progress", found=len(companies), total=max_companies,
                  message=f"Загружаем карточку: {c['name'][:40]}…")
            c = _enrich(c)
            _sleep(0.5, 1.5)   # polite delay after detail request

            companies.append(c)
            if len(companies) >= max_companies:
                break

        skip += 10
        page += 1
        _sleep(1.0, 2.5)   # delay between search pages

    return companies

# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

COLUMNS = [
    ("Название",          30),
    ("Телефон",           22),
    ("Сайт",              28),
    ("Социальные сети",   30),
    ("Адрес",             40),
    ("Широта",            14),
    ("Долгота",           14),
    ("Категория",         25),
    ("Рейтинг",           10),
    ("Кол-во отзывов",    16),
    ("Есть сайт",         12),
]

HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
HEADER_FONT = Font(bold=True)


def _write_header(ws, row: int = 1):
    for ci, (title, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row, ci, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_row(ws, ri: int, c: dict):
    vals = [
        c["name"], c["phone"], c["site"], c["social"], c["address"],
        c["lat"]  if isinstance(c["lat"],  float) else None,
        c["lon"]  if isinstance(c["lon"],  float) else None,
        c["category"],
        c["rating"]  if isinstance(c["rating"],  float) else None,
        c["reviews"],
        c["has_site"],
    ]
    for ci, v in enumerate(vals, 1):
        cell = ws.cell(ri, ci, v)
        if ci in (6, 7) and v is not None:
            cell.number_format = "0.000000"
        elif ci == 9 and v is not None:
            cell.number_format = "0.0"
        elif ci == 10:
            cell.number_format = "0"


def _set_widths(ws):
    for ci, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w


def _add_filter(ws, n: int):
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"


def _add_dv(ws, n: int):
    dv = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=False)
    dv.sqref = MultiCellRange(f"K2:K{n + 1}")
    ws.add_data_validation(dv)


def export_excel(companies: list[dict], query: str) -> str:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Компании"
    _write_header(ws1)
    for i, c in enumerate(companies, 2):
        _write_row(ws1, i, c)
    _set_widths(ws1)
    _add_filter(ws1, len(companies))
    _add_dv(ws1, len(companies))

    ws2 = wb.create_sheet("Без сайта")
    _write_header(ws2)
    no_site = [c for c in companies if c["has_site"] == "Нет"]
    for i, c in enumerate(no_site, 2):
        _write_row(ws2, i, c)
    _set_widths(ws2)
    _add_filter(ws2, len(no_site))
    _add_dv(ws2, len(no_site))

    safe = re.sub(r"[^\w\-а-яА-Я]", "_", query)[:40]
    filename = f"yandex_maps_{safe}_{date.today()}.xlsx"
    wb.save(filename)
    return filename

# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------

def print_stats(companies: list[dict], filename: str):
    total      = len(companies)
    with_site  = sum(1 for c in companies if c["has_site"] == "Да")
    with_phone = sum(1 for c in companies if c["phone"]    != "—")
    ratings    = [c["rating"] for c in companies if isinstance(c["rating"], float)]
    avg        = round(sum(ratings) / len(ratings), 1) if ratings else 0.0
    pct        = lambda n: f"{round(n / total * 100)}%" if total else "0%"

    print()
    print("=" * 46)
    print(f"  Файл: {filename}")
    print("=" * 46)
    print(f"  Всего          : {total}")
    print(f"  С сайтом       : {with_site} ({pct(with_site)})")
    print(f"  Без сайта      : {total - with_site} ({pct(total - with_site)})")
    print(f"  С телефоном    : {with_phone} ({pct(with_phone)})")
    print(f"  Средний рейтинг: {avg}")
    print("=" * 46)

# ---------------------------------------------------------------------------
# CLI entry point
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(
        description="Парсер Яндекс Карт → Excel"
    )
    ap.add_argument("query", help='Запрос, например "стоматология Санкт-Петербург"')
    ap.add_argument("-n", "--max", type=int, default=100, metavar="N",
                    help="Лимит компаний (по умолчанию 100)")
    args = ap.parse_args()

    log.info("Запрос: «%s» | Лимит: %d", args.query, args.max)
    companies = collect(args.query, args.max)

    if not companies:
        log.error("Ничего не найдено.")
        sys.exit(1)

    filename = export_excel(companies, args.query)
    print_stats(companies, filename)


if __name__ == "__main__":
    main()
