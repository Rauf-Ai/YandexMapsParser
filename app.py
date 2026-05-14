"""Flask web interface — Yandex Maps Parser"""

import json
import logging
import os
import queue
import re
import threading
import uuid
from datetime import date, datetime
from pathlib import Path

from flask import Flask, Response, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

from yandex_parser import (
    ALL_FIELD_DEFS, DEFAULT_FIELDS,
    apply_filters, collect, export_excel,
    _build_columns,
)

app = Flask(__name__)
DOWNLOADS_DIR = Path(os.environ.get("DOWNLOADS_DIR", "downloads"))
DOWNLOADS_DIR.mkdir(parents=True, exist_ok=True)
HISTORY_FILE  = Path(os.environ.get("DOWNLOADS_DIR", "downloads")) / "history.json"

JOBS: dict[str, dict] = {}
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# History helpers
# ---------------------------------------------------------------------------
def _load_history() -> list:
    if HISTORY_FILE.exists():
        try:
            return json.loads(HISTORY_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return []

def _save_history(entries: list):
    try:
        HISTORY_FILE.write_text(
            json.dumps(entries, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        log.warning("Could not save history: %s", exc)

# ---------------------------------------------------------------------------
# Excel writer (saves to explicit path, same logic as yandex_parser)
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
HEADER_FONT = Font(bold=True)


def _export_to_path(companies: list, filepath: Path,
                    selected_fields: list[str] | None = None):
    cols = _build_columns(selected_fields or DEFAULT_FIELDS)

    def _write_sheet(ws, rows):
        for ci, (_, header, _) in enumerate(cols, 1):
            cell = ws.cell(1, ci, header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for ri, c in enumerate(rows, 2):
            for ci, (key, _, _) in enumerate(cols, 1):
                val = c.get(key)
                if key in ("lat", "lon") and not isinstance(val, float):
                    val = None
                if key == "rating" and not isinstance(val, float):
                    val = None
                cell = ws.cell(ri, ci, val)
                if key in ("lat", "lon") and val is not None:
                    cell.number_format = "0.000000"
                elif key == "rating" and val is not None:
                    cell.number_format = "0.0"
                elif key == "reviews":
                    cell.number_format = "0"
        for ci, (_, _, width) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(ci)].width = width
        n = len(rows)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
        if n > 0:
            for ci, (key, _, _) in enumerate(cols, 1):
                if key == "has_site":
                    dv = DataValidation(
                        type="list", formula1='"Да,Нет"', allow_blank=False)
                    dv.sqref = MultiCellRange(
                        f"{get_column_letter(ci)}2:{get_column_letter(ci)}{n+1}")
                    ws.add_data_validation(dv)
                    break

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Компании"
    _write_sheet(ws1, companies)
    ws2 = wb.create_sheet("Без сайта")
    _write_sheet(ws2, [c for c in companies if c.get("has_site") == "Нет"])
    wb.save(filepath)

# ---------------------------------------------------------------------------
# Background job
# ---------------------------------------------------------------------------
def _run_job(job_id: str, query: str, max_companies: int,
             filters: dict, selected_fields: list[str]):
    job = JOBS[job_id]
    q: queue.Queue = job["queue"]

    def emit(event: str, **data):
        q.put({"event": event, "data": data})

    try:
        emit("start", message=f"Начинаем сбор: «{query}»")

        companies = collect(query, max_companies, emit=emit)

        if not companies:
            emit("error", message="Не удалось собрать ни одной компании.")
            job["status"] = "error"
            return

        # Apply filters
        filtered = apply_filters(
            companies,
            no_site=filters.get("no_site", False),
            no_social=filters.get("no_social", False),
            no_phone=filters.get("no_phone", False),
        )

        if not filtered:
            emit("error", message="После фильтрации не осталось компаний.")
            job["status"] = "error"
            return

        emit("progress", found=len(filtered), total=max_companies,
             message="Формируем Excel-файл…")

        safe     = re.sub(r"[^\w\-а-яА-Я]", "_", query)[:40]
        filename = f"yandex_maps_{safe}_{date.today()}.xlsx"
        filepath = DOWNLOADS_DIR / filename
        _export_to_path(filtered, filepath, selected_fields)

        total      = len(filtered)
        with_site  = sum(1 for c in filtered if c["has_site"] == "Да")
        with_phone = sum(1 for c in filtered if c["phone"]    != "—")
        with_social= sum(1 for c in filtered if c["social"]   != "—")
        ratings    = [c["rating"] for c in filtered
                      if isinstance(c["rating"], float)]
        avg        = round(sum(ratings) / len(ratings), 1) if ratings else 0.0

        job["status"]   = "done"
        job["filename"] = filename

        # Save to history
        hist  = _load_history()
        entry = {
            "id":       job_id,
            "query":    query,
            "date":     date.today().isoformat(),
            "time":     datetime.now().strftime("%H:%M"),
            "total":    total,
            "with_site":with_site,
            "filename": filename,
            "filters":  filters,
        }
        hist.insert(0, entry)
        _save_history(hist[:50])   # keep last 50

        emit("done",
             filename=filename,
             total=total,
             with_site=with_site,
             without_site=total - with_site,
             with_phone=with_phone,
             with_social=with_social,
             avg_rating=avg,
             pct_site=round(with_site  / total * 100) if total else 0,
             pct_phone=round(with_phone / total * 100) if total else 0,
             preview=[{
                 "name":    c["name"],
                 "phone":   c["phone"],
                 "address": c["address"],
                 "site":    c["site"],
                 "social":  c["social"],
                 "rating":  c["rating"],
                 "reviews": c["reviews"],
                 "has_site":c["has_site"],
                 "map_url": c.get("map_url", ""),
             } for c in filtered[:8]])

    except Exception as exc:
        log.exception("Job %s crashed", job_id)
        job["status"] = "error"
        q.put({"event": "error", "data": {"message": str(exc)}})
    finally:
        q.put(None)

# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.get("/")
def index():
    return render_template("index.html",
                           field_defs=ALL_FIELD_DEFS,
                           default_fields=DEFAULT_FIELDS)


@app.post("/api/start")
def api_start():
    body = request.get_json(force=True)

    # Build query from category + city OR raw query
    category = (body.get("category") or "").strip()
    city     = (body.get("city")     or "").strip()
    raw      = (body.get("query")    or "").strip()

    if category and city:
        query = f"{category} {city}"
    elif category:
        query = category
    elif city and raw:
        query = f"{raw} {city}"
    else:
        query = raw

    if not query:
        return jsonify(error="Введите поисковый запрос или выберите категорию"), 400

    try:
        max_n = max(1, min(int(body.get("max", 50)), 500))
    except (TypeError, ValueError):
        max_n = 50

    filters = {
        "no_site":   bool(body.get("no_site")),
        "no_social": bool(body.get("no_social")),
        "no_phone":  bool(body.get("no_phone")),
    }

    # Validate selected fields
    all_keys = list(ALL_FIELD_DEFS.keys())
    sel = body.get("fields") or DEFAULT_FIELDS
    selected_fields = [f for f in sel if f in all_keys]
    if "name" not in selected_fields:
        selected_fields.insert(0, "name")

    job_id = str(uuid.uuid4())
    JOBS[job_id] = {
        "status": "running", "query": query,
        "filename": None, "error": None,
        "queue": queue.Queue(),
    }
    threading.Thread(
        target=_run_job,
        args=(job_id, query, max_n, filters, selected_fields),
        daemon=True,
    ).start()
    return jsonify(job_id=job_id, query=query)


@app.get("/api/stream/<job_id>")
def api_stream(job_id):
    if job_id not in JOBS:
        return jsonify(error="Not found"), 404

    def generate():
        q: queue.Queue = JOBS[job_id]["queue"]
        while True:
            msg = q.get()
            if msg is None:
                yield "event: close\ndata: {}\n\n"
                break
            data = json.dumps(msg["data"], ensure_ascii=False)
            yield f"event: {msg['event']}\ndata: {data}\n\n"

    return Response(generate(), mimetype="text/event-stream",
                    headers={"Cache-Control": "no-cache",
                             "X-Accel-Buffering": "no"})


@app.get("/api/download/<filename>")
def api_download(filename):
    safe = re.sub(r"[^\w\-а-яА-Я.]", "_", filename)
    path = DOWNLOADS_DIR / safe
    if not path.exists():
        return jsonify(error="File not found"), 404
    return send_file(path, as_attachment=True, download_name=safe)


@app.get("/api/history")
def api_history():
    return jsonify(_load_history())


@app.delete("/api/history/<entry_id>")
def api_history_delete(entry_id):
    hist = [e for e in _load_history() if e.get("id") != entry_id]
    _save_history(hist)
    return jsonify(ok=True)


if __name__ == "__main__":
    app.run(debug=True, threaded=True, port=5000)
